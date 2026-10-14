import base64
from io import BytesIO
from datetime import datetime

from openpyxl import load_workbook

from odoo import fields, models, _
from odoo.exceptions import UserError


class CrmRequestImportWizard(models.TransientModel):
    _name = "crm.request.import.wizard"
    _description = "Import Customer Requests"

    opportunity_id = fields.Many2one(
        "crm.lead",
        string="Cơ hội",
        required=True,
        default=lambda self: self.env.context.get("active_id"),
    )

    file_data = fields.Binary(
        string="File Excel",
        required=True,
    )

    file_name = fields.Char(string="Tên file")

    def action_import(self):
        self.ensure_one()

        if hasattr(self.opportunity_id, "_is_request_editable") and not self.opportunity_id._is_request_editable():
            raise UserError(
                _("Chỉ được import khi cơ hội đang ở giai đoạn Mới.")
            )

        try:
            content = base64.b64decode(self.file_data)
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise UserError(
                _("Không thể đọc file Excel. Hãy chọn file .xlsx hợp lệ.")
            ) from exc

        if "Requests" not in workbook.sheetnames:
            raise UserError(
                _("File Excel phải có sheet tên Requests.")
            )

        sheet = workbook["Requests"]
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)  # Bỏ qua dòng tiêu đề

        products = self.env["product.product"]
        requests = self.env["crm.customer.request"]
        values_list = []

        for row_number, row in enumerate(rows, start=2):
            if not row or not any(value is not None for value in row):
                continue

            product_name = row[0]
            qty = row[1]
            request_date = row[2]
            description = row[3] if len(row) > 3 else False

            if not product_name or qty is None:
                raise UserError(
                    _("Thiếu sản phẩm hoặc số lượng tại dòng %s.")
                    % row_number
                )

            product = products.search(
                [("name", "=", str(product_name).strip())],
                limit=2,
            )

            if len(product) != 1:
                raise UserError(
                    _("Không tìm thấy duy nhất sản phẩm '%s' tại dòng %s.")
                    % (product_name, row_number)
                )

            try:
                qty = float(qty)
            except (TypeError, ValueError) as exc:
                raise UserError(
                    _("Số lượng không hợp lệ tại dòng %s.") % row_number
                ) from exc

            if qty <= 0:
                raise UserError(
                    _("Số lượng phải lớn hơn 0 tại dòng %s.")
                    % row_number
                )

            # Xử lý linh hoạt định dạng Ngày
            parsed_date = False
            if request_date:
                if isinstance(request_date, datetime):
                    parsed_date = request_date.date()
                elif hasattr(request_date, "date"):
                    parsed_date = request_date.date()
                else:
                    date_str = str(request_date).strip()
                    # Thử ép kiểu theo các định dạng ngày phổ biến
                    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
                        try:
                            parsed_date = datetime.strptime(date_str, fmt).date()
                            break
                        except ValueError:
                            continue

                    if not parsed_date:
                        try:
                            parsed_date = fields.Date.to_date(date_str)
                        except Exception as exc:
                            raise UserError(
                                _("Ngày không hợp lệ tại dòng %s.") % row_number
                            ) from exc

            if not parsed_date:
                parsed_date = fields.Date.context_today(self)

            values_list.append({
                "opportunity_id": self.opportunity_id.id,
                "product_id": product.id,
                "qty": qty,
                "date": parsed_date,
                "description": str(description) if description else False,
            })

        if not values_list:
            raise UserError(_("File Excel không có dòng nhu cầu hợp lệ."))

        requests.create(values_list)

        return {"type": "ir.actions.act_window_close"}